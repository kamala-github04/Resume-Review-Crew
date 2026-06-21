"""
Utilities for extracting plain text from uploaded resume files (PDF or DOCX).
"""
from io import BytesIO

from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file given as raw bytes."""
    reader = PdfReader(BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file given as raw bytes."""
    document = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Resumes sometimes use table-based layouts; pull that text too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()


def extract_resume_text(uploaded_file) -> str:
    """
    Extract text from a Streamlit UploadedFile object.

    Supports .pdf and .docx. Raises ValueError on unsupported file types
    or if no text could be extracted (e.g. a scanned image-only PDF).
    """
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if filename.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file type. Please upload a .pdf or .docx file.")

    if not text:
        raise ValueError(
            "Could not extract any text from this file. It may be a scanned "
            "image-based PDF without selectable text — try exporting a text-based "
            "version instead."
        )

    return text
